# SPDX-License-Identifier: AGPL-3.0-or-later

"""Isolated Phase 1 browser-test app backed by a fake Moodle connector."""

import copy
import hashlib
import json
import os
import re
import shutil
import sqlite3
import tempfile
from dataclasses import asdict
from datetime import UTC, datetime, time, timedelta
from email.header import decode_header, make_header
from io import BytesIO
from pathlib import Path
from threading import Event
from types import SimpleNamespace
from uuid import uuid4

import httpx
import pymupdf as fitz
from alembic.config import Config
from apscheduler.events import EVENT_JOB_EXECUTED
from fastapi.responses import Response
from pptx import Presentation
from sqlalchemy import delete, select, text

from alembic import command

_sandbox = tempfile.TemporaryDirectory(prefix="marc-uni-os-e2e-")
_data_dir = Path(_sandbox.name)
(_data_dir / "drop").mkdir()
(_data_dir / "drop" / "Eigene_Frage.txt").write_text(
    "Warum ist eine Testmatrix invertierbar?", encoding="utf-8",
)
_live_db_path = _data_dir / "private" / "test.sqlite3"

from app import security as security_module
from app.ai import proposals_service as ai_api_module
from app.ai import provider_repository as ai_provider_repository
from app.ai import provider_service as ai_provider_service
from app.ai.adapter import AdapterResult
from app.config import Settings
from app.connectors.base import ConnectorAuth
from app.db.models import (
    AnkiCard,
    AnkiReview,
    Assignment,
    Course,
    Deadline,
    DeadlineProposal,
    Exam,
    ExerciseProposalQueue,
    Note,
    NoteVersion,
    Notification,
    Resource,
    ResourceVersion,
    StudyBlockSuggestion,
    StudyCycle,
    SubmissionAttempt,
    SyncRun,
    utc_now,
)
from app.files.atomic import ensure_directory, write_once
from app.health import doctor as doctor_module
from app.health.doctor import Diagnostic
from app.host.lease import HostCoordinator
from app.host.snapshots import SnapshotManifest, create_snapshot
from app.learning.topics import resolve_topic
from app.main import create_app
from app.notes.store import create_note
from app.notify.service import Notifier
from app.submissions import api as submission_api
from app.sync.engine import SyncEngine
from app.sync.scheduler import SyncScheduler
from tests.sample_files import sample_files
from tests.test_moodle_sync import (
    FixtureMoodle,
)

# Browser fixtures reset and seed the fake API directly. Session/fencing gates
# are covered separately with the production setting in backend tests.
settings = Settings(
    _env_file=None,  # type: ignore[call-arg]  # pydantic-settings runtime option
    data_dir=_data_dir,
    local_dir=_data_dir / "private",
    database_url=f"sqlite:///{_live_db_path}",
    device_name="Test-Mac",
    institution_slug="hs-coburg",
    moodle_enabled=True,
    moodle_base_url="https://moodle.example.invalid",
    moodle_token="fixture-token",
    syncthing_api_url="http://127.0.0.1:8177",
    syncthing_folder_id="fixture-uni-data",
    syncthing_api_key="fixture-syncthing-key",  # pragma: allowlist secret
    tailscale_ip="100.101.102.103",
    server_mode=False,
    demo_mode=False,
    public_server_url="",
    public_periods_enabled=False,
    api_auth_required=False,
    ntfy_url="",
    discord_webhook_url="",
    host_health_url="",
    restic_repository="",
    restic_password="",
    mcp_enabled=False,
    mcp_token="",
    remote_mcp_public_url="",
    remote_mcp_approval_pin="",
    anki_agent_token="",
    github_token="",
    ai_adapter_enabled=False,
    ai_adapter_api_key="",
    ai_adapter_model="",
    ai_adapter_monthly_budget_tokens=None,
    ai_adapter_effort=None,
    connector_plugins="",
    allow_trusted_python_connectors=False,
    coding_runner_url="",
    coding_runner_token="",
    coding_runner_isolated=False,
    submission_tsa_url="",
    request_delay_seconds=0,
    tz="Europe/Berlin",
    allowed_hosts="127.0.0.1,localhost,[::1],testserver",
    serve_react=os.environ.get("SERVE_REACT", "").casefold() == "true",
    react_dist_dir=Path(os.environ.get("UNIOS_TEST_REACT_DIST", str(Path(__file__).resolve().parents[2] / "frontend" / "dist"))),
)
# This explicitly named audit port is trusted only inside this isolated fake
# process. Keep production origins and all ordinary browser fixtures unchanged.
_fake_preview_port = os.environ.get("UNIOS_FAKE_PREVIEW_PORT")
if _fake_preview_port is not None:
    if _fake_preview_port != "4194":
        raise ValueError("UNIOS_FAKE_PREVIEW_PORT unterstützt ausschließlich den QA-Port 4194")
    security_module.LOCAL_PREVIEW_ORIGINS = security_module.LOCAL_PREVIEW_ORIGINS | frozenset({
        "http://127.0.0.1:4194", "http://localhost:4194",
    })
syncthing_config = _data_dir / "syncthing-config.xml"
syncthing_config.write_text(
    f'<configuration><folder id="fixture-uni-data" path="{settings.data_dir}" '
    'type="sendreceive" /><device><address>tcp://peer.tailnet.ts.net:22000</address></device>'
    '<options><globalAnnounceEnabled>false</globalAnnounceEnabled>'
    '<relaysEnabled>false</relaysEnabled><natEnabled>false</natEnabled>'
    '<localAnnounceEnabled>false</localAnnounceEnabled>'
    '<listenAddress>tcp://100.101.102.103:22000</listenAddress></options></configuration>',
    encoding="utf-8",
)
doctor_module._syncthing_config_paths = lambda: [syncthing_config]
host = HostCoordinator(settings)
host.acquire_initial()
host.start()
config = Config(str(Path(__file__).resolve().parents[2] / "alembic.ini"))
config.set_main_option("script_location", str(Path(__file__).resolve().parents[2] / "alembic"))
config.attributes["settings"] = settings
config.attributes["host_coordinator"] = host
command.upgrade(config, "head")
_original_select_adapter = ai_api_module.select_adapter
fixture = FixtureMoodle(settings)
fixture.responses["core_group_get_course_user_groups"] = {"groups": [{"name": "Müller"}]}
fixture.contents[101].append({"id": 302, "name": "Gruppe Schmidt", "modules": []})
pdf = fitz.open()
page = pdf.new_page()
page.insert_text((72, 72), "Schleifen und Python Beispiele")
page = pdf.new_page()
page.insert_text((72, 72), "Rekursionstiefe und Funktionsaufrufe")
pdf_data = pdf.tobytes()
pdf.close()
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
textbox = slide.shapes.add_textbox(0, 0, 6000000, 1000000)
textbox.text = "Mathematik Uebungsblatt und Aufgaben"
output = BytesIO()
slides.save(output)
pptx_data = output.getvalue()
fixture.files["/webservice/pluginfile.php/10/04_Schleifen.pdf"] = pdf_data
fixture.files["/webservice/pluginfile.php/20/Blatt_01.pptx"] = pptx_data
fixture.contents[101][0]["modules"][0]["contents"][0]["filesize"] = len(pdf_data)
fixture.contents[202][0]["modules"][0]["contents"][0]["filesize"] = len(pptx_data)
submission_file = b"Erfundene Abgabe"
fixture.files["/webservice/pluginfile.php/70/Abgabe.txt"] = submission_file
fake_submission_status = "draft"
uploaded_submission_name = None
uploaded_submission_content = None


class FakeSubmissionMoodle:
    def __init__(self, _settings):
        self.client = SimpleNamespace(close=lambda: None)

    def check_auth(self):
        return ConnectorAuth("42", frozenset({"mod_assign_get_submission_status", "mod_assign_get_assignments",
                                              "mod_assign_save_submission", "mod_assign_submit_for_grading"}))

    def call(self, function, **_parameters):
        global fake_submission_status
        if function == "mod_assign_get_assignments":
            listing = copy.deepcopy(fixture.responses[function])
            listing["courses"][0]["assignments"][0]["submissiondrafts"] = True
            listing["courses"][0]["assignments"][0]["requiresubmissionstatement"] = 1
            listing["courses"][0]["assignments"][0]["submissionstatement"] = "Erfundene Erklärung zur eigenen Arbeit."
            listing["courses"][0]["assignments"][0]["configs"] = [
                {"plugin": "file", "subtype": "assignsubmission", "name": "enabled", "value": "1"},
                {"plugin": "file", "subtype": "assignsubmission", "name": "maxfilesubmissions", "value": "3"},
                {"plugin": "file", "subtype": "assignsubmission", "name": "maxsubmissionsizebytes", "value": "1000000"},
                {"plugin": "file", "subtype": "assignsubmission", "name": "filetypeslist", "value": ".pdf,.zip"},
            ]
            return listing
        if function == "mod_assign_get_submission_status":
            filename = uploaded_submission_name or "Abgabe.txt"
            content = uploaded_submission_content or submission_file
            files = [{"filename": filename, "filesize": len(content),
                      "fileurl": f"https://moodle.example.invalid/webservice/pluginfile.php/70/{filename}"}]
            return {"lastattempt": {"graded": False, "submission": {
                "id": 70, "status": fake_submission_status, "timemodified": 1790000000,
                "plugins": [{"type": "file", "fileareas": [{"files": files}]}],
            }}}
        if function == "mod_assign_save_submission":
            if uploaded_submission_name is None:
                raise AssertionError("No fake draft uploaded")
            fake_submission_status = "draft"
            return {"warnings": []}
        if function == "mod_assign_submit_for_grading":
            fake_submission_status = "submitted"
            return {"warnings": []}
        raise AssertionError(function)

    def upload_draft(self, name, content, item_id=0):
        global uploaded_submission_name, uploaded_submission_content
        uploaded_submission_name = name
        uploaded_submission_content = content
        return 99

    def download(self, resource):
        if uploaded_submission_name and resource.name == uploaded_submission_name:
            return uploaded_submission_content
        return fixture.connector.download(resource)


submission_api.MoodleClient = FakeSubmissionMoodle  # type: ignore[misc,assignment]  # swaps the whole connector class for this fake app process, not per-test monkeypatching
import app.main as main_module
from tests.fake_servers.personal_sources import FakeMailService, FakeTimetableService

main_module.MailService = FakeMailService  # type: ignore[misc,assignment]
main_module.TimetableService = FakeTimetableService  # type: ignore[misc,assignment]
from app.mail.forwarding import MailForwardingService


def _fake_forwarding(settings, factory, coordinator):
    global fake_forwarding
    fake_forwarding = MailForwardingService(settings.model_copy(update={"server_mode": True}), factory, coordinator)
    return fake_forwarding


main_module.MailForwardingService = _fake_forwarding  # type: ignore[misc,assignment]
from app.ai import connection as mcp_connection_module
from tests.test_mcp_connection import MemoryVault as McpMemoryVault

mcp_connection_module.Vault = McpMemoryVault  # type: ignore[misc,assignment]
app = create_app(settings, connectors=[fixture.connector], start_scheduler=False, coordinator=host)
from tests.fake_servers.coding_runner import BrowserCodingRunner

app.state.coding_service.backend = BrowserCodingRunner()
result = app.state.sync_engine.run()
if result is None or result.status != "success":
    raise RuntimeError("Fake-Moodle-Start-Sync fehlgeschlagen")
_startup_snapshot = create_snapshot(host)
original_credential_file_check = doctor_module._credential_file
syncthing_conflict_file = settings.data_dir / "host" / "lease.sync-conflict-20260925-073000-TEST.json"

# Baseline state captured once, right after startup finished its own sync.
# "/api/__test__/reset" restores exactly this snapshot so specs that run later
# in the same Playwright worker don't inherit state left behind by earlier
# specs (see the reset route near the end of this module).
#
# The database half is a plain file copy of the snapshot create_snapshot()
# already verified above, kept outside data/host/snapshots/ so the snapshot
# retention pruning (settings.snapshot_keep) never deletes it. Restoring by
# copying the whole file back (rather than deleting ORM rows and re-running
# the sync) keeps the DB byte-consistent with the untouched, content-addressed
# files already on disk under data/uni/ — re-running the sync against an
# emptied database re-downloads content whose immutable per-version sidecar
# JSON already exists on disk with a different timestamp, which the
# write-once file store correctly refuses to overwrite.
_reset_baseline_db = _data_dir / "private" / "reset-baseline.sqlite3"
shutil.copy2(settings.data_dir / "host" / "snapshots" / _startup_snapshot.filename, _reset_baseline_db)
_seed_drop_question = (_data_dir / "drop" / "Eigene_Frage.txt").read_text(encoding="utf-8")
_fixture_baseline = {
    "contents": copy.deepcopy(fixture.contents),
    "responses": copy.deepcopy(fixture.responses),
    "files": copy.deepcopy(fixture.files),
    "download_count": fixture.download_count,
}
# DATA_DIR subfolders that tests write into; reset() clears them and restores
# only the one file that was seeded before the app ever started.
_RESET_DATA_SUBFOLDERS = ("annotations", "drop", "ai/inbox", "ai/outbox", "submissions", "notes",
                        "host/conflicts", "backup-status", "host/snapshots", "uni/manual-card-fixtures",
                        "note-media", "voice")
_snapshot_baseline = {path.name: path.read_bytes() for path in (settings.data_dir / "host/snapshots").iterdir()
                      if path.is_file()}
_lease_baseline = host.lease_path.read_bytes()


@app.post('/api/__test__/timetable-slots')
def timetable_slots_for_browser_test():
    from app.calendar.service import TimetableService
    service = TimetableService(settings, app.state.factory, host)
    rows = [{'source_key': f'fictional-{hour}', 'title': 'English Workshop · Testbau A-12',
             'starts_at': datetime(2026, 10, 5, hour, minute, tzinfo=UTC),
             'ends_at': datetime(2026, 10, 5, hour + 1, minute, tzinfo=UTC),
             'course_id': None, 'url': None} for hour, minute in [(8, 0), (9, 45)]]
    service.repository.replace(service.source, rows, host)
    service.store.save({'institution': settings.institution_slug, 'subscribed': True})
    service.selection.save(set())
    return {'status': 'ok'}


@app.post('/api/__test__/timetable-series')
def timetable_series_for_browser_test():
    from zoneinfo import ZoneInfo

    from app.calendar.service import TimetableService
    service = TimetableService(settings, app.state.factory, host)
    zone = ZoneInfo('Europe/Berlin')
    rows = []
    for day in ['2026-10-06', '2026-10-13', '2026-10-27', '2026-10-08']:
        start = datetime.fromisoformat(day + 'T14:15').replace(tzinfo=zone)
        rows.append({'source_key': 'series-' + day, 'title': 'Fiktive Methoden · Testbau A-12',
                     'starts_at': start.astimezone(UTC), 'ends_at': start.replace(hour=15, minute=30).astimezone(UTC),
                     'course_id': None, 'url': None})
    service.repository.replace(service.source, rows, host)
    service.store.save({'institution': settings.institution_slug, 'subscribed': True})
    service.selection.save(set())
    return {'status': 'ok'}


@app.post("/api/__test__/forwarded-mail")
def forwarded_mail_for_browser_test():
    from email.message import EmailMessage
    fake_forwarding.configuration.save({"address": "inbox@mail.example.test", "ready": True})
    fake_forwarding.state.save({"enabled": True})
    for path in (fake_forwarding.mailbox.parent, fake_forwarding.mailbox,
                 fake_forwarding.mailbox / "new", fake_forwarding.mailbox / "cur"):
        path.mkdir(exist_ok=True)
        path.chmod(0o700)
    mail = EmailMessage()
    mail["From"] = "office@example.test"
    mail["To"] = "inbox@mail.example.test"
    mail["Subject"] = "Erfundene Weiterleitungsfrist"
    due = datetime.now(UTC) + timedelta(days=8)
    mail.set_content(f"Der Abgabeschluss ist am {due:%d.%m.%Y}.")
    file = fake_forwarding.mailbox / "new" / uuid4().hex
    file.write_bytes(mail.as_bytes())
    file.chmod(0o600)
    return {"status": "ok"}
@app.get("/api/__test__/long-reader.pdf")
def long_reader_for_browser_test(pages: int = 20):
    with fitz.open() as document:
        for number in range(1, min(500, max(1, pages)) + 1):
            page = document.new_page(width=960 if number % 2 == 0 else 595,
                                     height=540 if number % 2 == 0 else 842)
            page.insert_text((40, 50), f"Erfundene Leseseite {number}")
            page.insert_text((40, page.rect.height - 35), f"Seitenende {number}")
        return Response(document.tobytes(), media_type="application/pdf")


def _study_fixture_directories() -> tuple[Path, Path]:
    root = _data_dir.resolve()
    folders = (settings.local_dir / "anki", settings.data_dir / "anki")
    for folder in folders:
        if folder.is_symlink() or not folder.resolve().is_relative_to(root):
            raise RuntimeError("Anki-Testzustand liegt außerhalb des temporären Fake-Verzeichnisses")
    return folders


def _reset_study_fixture() -> None:
    folders = _study_fixture_directories()
    if not any(folder.exists() for folder in folders):
        return
    from app.anki.collection import _lock, close_study_collection

    # Finish the bounded native undo session under the production collection
    # lock before removing these verified, synthetic fixture-only directories.
    with _lock:
        close_study_collection(settings)
        for folder in _study_fixture_directories():
            if folder.exists():
                shutil.rmtree(folder)


@app.post("/api/__test__/study-deck-scenario")
def seed_study_deck_scenario():
    """Seed real Anki queues beyond the old global 100-card boundary."""
    from app.anki.collection import ensure_notetype, open_collection, write_study_mode

    host.assert_writer()
    _reset_study_fixture()
    with open_collection(settings) as collection:
        model = ensure_notetype(collection, "basic")
        names = {"parent": "Uni", "deck_a": "Uni::A", "deck_b": "Uni::B",
                 "deck_bb": "Uni::BB", "deck_c": "Uni::C"}
        decks: dict[str, dict] = {}
        for key, name in names.items():
            identifier = collection.decks.id(name)
            assert identifier is not None
            config = collection.decks.config_dict_for_deck_id(identifier)
            config["new"]["perDay"] = 1000
            config["rev"]["perDay"] = 1000
            config["new"]["order"] = 1
            config["newGatherPriority"] = 0
            config["newSortOrder"] = 0
            config["newMix"] = 0
            collection.decks.update_config(config)
            decks[key] = {"id": int(identifier), "name": name, "card_ids": [], "fronts": []}
        position = 0
        for key, count, course in (("deck_a", 101, "Programmieren 1"), ("deck_b", 1, "Mathematik"),
                                   ("deck_bb", 1, "Mathematik"), ("deck_c", 1, "Programmieren 1")):
            for index in range(1, count + 1):
                position += 1
                letter = key.removeprefix("deck_").upper()
                front = f"{letter}-Karte {index:03}: Was erklärt das erfundene Beispiel?"
                note = collection.new_note(model)
                note["Front"] = front
                note["Back"] = f"Die synthetische Erklärung für {letter}-{index:03}."
                for field, value in {
                    "Course": course, "Topic": "Deck-Scope", "Semester": "Semester_1",
                    "SourceFile": "Erfundene_Deckbeispiele.pdf", "SourcePage": "1",
                    "CreatedAt": datetime.now(UTC).isoformat(), "UniOSId": f"scope-fixture-{position}",
                }.items():
                    note[field] = value
                note.tags = ["uni::synthetisch::deck-scope", "semester::1"]
                collection.add_note(note, decks[key]["id"])
                decks[key]["card_ids"].extend(int(card.id) for card in note.cards())
                decks[key]["fronts"].append(front)
        collection.decks.select(decks["parent"]["id"])
        queued = collection.sched.get_queued_cards(fetch_limit=101)
        first_ids = [int(item.card.id) for item in queued.cards]
        if len(first_ids) != 101 or not set(first_ids).issubset(decks["deck_a"]["card_ids"]):
            raise RuntimeError("Anki-Testkonfiguration bildet die beabsichtigte Queuegrenze nicht ab")
        collection.decks.select(decks["deck_c"]["id"])
        selected = int(collection.decks.selected())
    write_study_mode(settings, "integrated")
    return {**decks, "selected_deck_id": selected, "first_parent_card_id": first_ids[0],
            "native_parent_counts": {"new": queued.new_count, "learning": queued.learning_count,
                                     "review": queued.review_count}, "total_cards": 104}


@app.get("/api/__test__/study-deck-scenario")
def inspect_study_deck_scenario():
    """Inspect only the synthetic collection, including exactly-once replay."""
    from app.anki.collection import collection_path, open_collection

    _study_fixture_directories()
    if not collection_path(settings).is_file():
        return {"present": False}
    with open_collection(settings) as collection:
        assert collection.db is not None
        receipts = []
        if collection.db.scalar("SELECT count(*) FROM sqlite_master WHERE name = 'unios_review_receipts'"):
            receipts = [{"review_id": row[0], "card_id": row[1], "rating": row[2]} for row in collection.db.all(
                "SELECT review_id, card_id, rating FROM unios_review_receipts ORDER BY review_id")]
        result = {"present": True, "selected_deck_id": int(collection.decks.selected()),
                "cards": len(collection.find_cards("")),
                "reviews": int(collection.db.scalar("SELECT count(*) FROM revlog") or 0),
                "receipts": receipts,
                "revlog": [{"id": row[0], "card_id": row[1], "rating": row[2]} for row in collection.db.all(
                    "SELECT id, cid, ease FROM revlog ORDER BY id")]}
    with app.state.factory() as session:
        result["sql_reviews"] = [{"id": review.id, "card_id": review.card_id,
                                  "rating": review.rating, "review_id": review.review_id,
                                  "reviewed_at": review.reviewed_at.replace(tzinfo=UTC).isoformat()}
                                 for review in session.scalars(select(AnkiReview).order_by(AnkiReview.id)).all()]
        result["sql_cards"] = [{"id": card.id, "queue_id": card.queue_id, "note_id": card.note_id}
                               for card in session.scalars(select(AnkiCard).order_by(AnkiCard.id)).all()]
    return result


def _disarm_review_fault() -> None:
    patcher = getattr(app.state, "review_fault_patch", None)
    if patcher is not None:
        patcher.stop()
        app.state.review_fault_patch = None


@app.post("/api/__test__/study-review-fault")
def fail_review_mirror_once():
    """Fail only the synthetic SQL mirror after the real native receipt commit."""
    from unittest.mock import patch

    from app.anki.study_review_repository import ReviewMirrorRepository

    host.assert_writer()
    _disarm_review_fault()

    def fail_once(*_args, **_kwargs):
        _disarm_review_fault()
        raise OSError("Synthetischer Schreibausfall nach nativem Bewertungs-Commit")

    patcher = patch.object(ReviewMirrorRepository, "add_review", autospec=True, side_effect=fail_once)
    app.state.review_fault_patch = patcher
    patcher.start()
    return {"status": "armed", "synthetic_only": True}


@app.post("/api/__test__/study-note-rebind")
def rebind_synthetic_native_note(queue_id: int):
    """Simulate import-assigned native note IDs while preserving the stable GUID."""
    from app.anki.collection import open_collection

    host.assert_writer()
    _study_fixture_directories()
    with app.state.factory() as session:
        mapped = session.scalar(select(AnkiCard).where(AnkiCard.queue_id == queue_id))
        if mapped is None:
            raise ValueError("Synthetische Karte ist nicht zugeordnet")
        old_id = int(mapped.note_id)
        sql_card_id = mapped.id
    with open_collection(settings) as collection:
        assert collection.db is not None
        new_id = old_id + 1_000_000
        if collection.db.scalar("SELECT count(*) FROM notes WHERE id = ?", new_id):
            raise ValueError("Synthetische Zielkennung ist bereits belegt")

        def apply():
            collection.db.execute("UPDATE notes SET id = ? WHERE id = ?", new_id, old_id)
            collection.db.execute("UPDATE cards SET nid = ? WHERE nid = ?", new_id, old_id)

        collection.db.transact(apply)
    return {"old_note_id": str(old_id), "new_note_id": str(new_id), "sql_card_id": sql_card_id}


@app.post("/api/__test__/study-native-card-remove")
def remove_synthetic_native_card(card_id: int):
    """Official SDK removal only inside this isolated temporary browser fixture."""
    from anki.collection import CardId

    from app.anki.collection import open_collection

    host.assert_writer()
    _study_fixture_directories()
    with open_collection(settings) as collection:
        note_id = collection.get_card(CardId(card_id)).nid
        collection.remove_notes([note_id])
    return {"removed_card_id": card_id}


@app.post("/api/__test__/manual-card-scenario")
def seed_manual_card_scenario(mode: str = "integrated"):
    """Two actual 120-page PDFs with the same passage on different pages."""
    from app.anki.collection import write_study_mode
    from app.material.service import create, display

    host.assert_writer()
    if mode not in {"integrated", "external"}:
        raise ValueError("Ungültiger synthetischer Lernmodus")
    _reset_study_fixture()
    _study_fixture_directories()
    write_study_mode(settings, mode)
    quote = "Das synthetische Prinzip verbindet Begriffe mit ihrer urspruenglichen Quelle."
    marker = uuid4().hex
    folder = _data_dir / "uni" / "manual-card-fixtures"
    folder.mkdir(parents=True, exist_ok=True)
    versions = []
    with app.state.factory() as session:
        course = session.get(Course, 1)
        assert course is not None
        resource = Resource(source="fixture", source_key=f"manual-{marker}", course_id=course.id,
                            kind="file", name="Synthetisches_Lernskript.pdf", url="", group_category="general")
        session.add(resource)
        session.flush()
        for number, quote_page in ((1, 12), (2, 17)):
            with fitz.open() as document:
                for page_number in range(1, 121):
                    page = document.new_page()
                    page.insert_text((70, 90), f"Erfundenes Lernskript V{number} Seite {page_number}")
                    page.insert_text((70, 135), quote if page_number == quote_page else
                                     f"Synthetisches Kapitel {page_number}: Begriffe und Beispiele.")
                content = document.tobytes()
            path = folder / f"{marker}-v{number}.pdf"
            path.write_bytes(content)
            version = ResourceVersion(resource_id=resource.id, version=number,
                                      sha256=hashlib.sha256(content).hexdigest(),
                                      relative_path=path.relative_to(_data_dir).as_posix(),
                                      size=len(content), mime_type="application/pdf")
            session.add(version)
            session.flush()
            versions.append(version)
        annotation = create(session, _data_dir, versions[0], "question", quote, 12,
                            "Wie bleiben Begriffe mit ihrer Quelle verbunden?")
        session.commit()
        displayed = display(session, _data_dir, versions[0], annotation)
        return {"version_id": versions[0].id, "next_version_id": versions[1].id,
                "resource_id": resource.id, "course_id": course.id,
                "deck_name": f"Uni::{course.name}", "quote": quote, "page": 12,
                "annotation": displayed}


@app.post("/api/__test__/focus-workspace")
def seed_focus_workspace():
    """Accepted work beside 201 unrelated, untouched proposals; no live study data."""
    from app.ai.exercises import ExerciseProposal

    source = seed_manual_card_scenario()
    now = utc_now()
    with app.state.factory() as session:
        course = session.get(Course, source["course_id"])
        assert course is not None
        other_course = session.scalar(select(Course).where(Course.id != course.id).order_by(Course.id))
        if other_course is None:
            raise RuntimeError("Der synthetische Fokusablauf benötigt zwei Kurse")
        course.name = "Analysis – Fokusbeispiel"
        other_course.name = "Programmierung – anderes Fach"
        topic = resolve_topic(session, course.id, "Grenzwerte")
        exam_at = now + timedelta(days=5)
        exam = Exam(course_id=course.id, source_key=f"focus-workspace-{uuid4().hex}",
                    title="Analysis – Beispielprüfung", exam_at=exam_at)
        session.add(exam)
        session.flush()
        block = StudyBlockSuggestion(exam_id=exam.id, topic_id=topic.id, duration_minutes=35,
                                     starts_at=now, status="accepted", created_at=now)
        cycle = StudyCycle(exam_id=exam.id, phase="practice", starts_at=now,
                           ends_at=exam_at - timedelta(days=1), exam_at_snapshot=exam_at,
                           topic_ids=[topic.id], status="accepted", review_goal=1, attempt_goal=1)
        session.add_all([block, cycle])
        session.flush()
        note = create_note(session, settings, course_id=course.id, resource_id=source["resource_id"],
                           title="Meine Grenzwert-Erklärung", body="# Grenzwerte\n\nMeine eigene Begründung.")
        foreign_note = create_note(session, settings, course_id=other_course.id, resource_id=None,
                                   title="Notiz aus dem anderen Fach", body="Erfundene Programmiernotiz.")
        original_version = session.get(ResourceVersion, source["version_id"])
        assert original_version is not None
        foreign_resource = Resource(source="fixture", source_key=f"focus-foreign-{uuid4().hex}",
                                    course_id=other_course.id, kind="file", name="Fremdes_Skript.pdf",
                                    url="", group_category="general")
        session.add(foreign_resource)
        session.flush()
        foreign_version = ResourceVersion(resource_id=foreign_resource.id, version=1,
                                          relative_path=original_version.relative_path, sha256=original_version.sha256,
                                          size=original_version.size, mime_type="application/pdf")
        session.add(foreign_version)
        session.flush()
        proposal = ExerciseProposal(question="Warum genügt bei 3x + 1 die Wahl δ = ε / 3?",
                                    topic=topic.name, difficulty="medium", sample_solution="Die Steigung ist 3.",
                                    solution_steps="Der Ergebnisabstand ist dreimal so groß.", source_page=12)
        exercise = ExerciseProposalQueue(resource_id=source["resource_id"], status="approved",
                                         payload=proposal.model_dump_json())
        session.add(exercise)
        session.flush()
        other_resource = session.scalar(select(Resource).where(Resource.course_id == other_course.id,
                                                               Resource.group_category != "other_group"))
        if other_resource is None:
            raise RuntimeError("Das synthetische andere Fach benötigt eine Quelle")
        for index in range(201):
            unrelated = proposal.model_copy(update={"question": f"Fremde Programmieraufgabe {index + 1}",
                                                     "topic": "Fremder Programmierstoff"})
            session.add(ExerciseProposalQueue(resource_id=other_resource.id, status="approved",
                                              payload=unrelated.model_dump_json()))
        session.commit()
        return {**source, "block_id": block.id, "cycle_id": cycle.id, "topic_id": topic.id,
                "course_name": course.name, "deck_name": f"Uni::{course.name}",
                "exam_at": exam_at.isoformat(), "exercise_id": exercise.id,
                "note_id": note.id, "foreign_note_id": foreign_note.id,
                "foreign_version_id": foreign_version.id,
                "foreign_course_id": other_course.id}


@app.post("/api/__test__/focus-course-rename")
def rename_focus_course(course_id: int):
    host.assert_writer()
    with app.state.factory() as session:
        course = session.get(Course, course_id)
        if course is None:
            raise RuntimeError("Synthetischer Fokus-Kurs fehlt")
        course.name = "Analysis – aktueller Kursname"
        session.commit()
    return {"name": "Analysis – aktueller Kursname"}


@app.post("/api/__test__/legacy-note")
def legacy_note_for_browser_test(course_id: int):
    """A note stored the way releases before the block format wrote it: one Markdown revision file."""
    with app.state.factory() as session:
        course = session.get(Course, course_id)
        if course is None:
            raise RuntimeError("Synthetischer Kurs fehlt")
        note_id = str(uuid4())
        body = "Historischer Markdown-Text"
        header = ("---\n"
                  f"unios_note_id: {json.dumps(note_id)}\n"
                  'title: "Frühere Frage"\n'
                  f"course: {json.dumps(course.name, ensure_ascii=False)}\n"
                  f"course_id: {course.id}\n"
                  "resource_id: null\n"
                  "revision: 1\n"
                  "---\n\n")
        folder = settings.data_dir / "notes" / f"legacy_{course.id}" / note_id
        folder.mkdir(parents=True)
        target = folder / "v0001_Fruehere_Frage_legacy01.md"
        content = (header + body).encode("utf-8")
        write_once(target, content)
        now = utc_now()
        session.add(Note(id=note_id, course_id=course.id, resource_id=None, title="Frühere Frage",
                         current_revision=1, created_at=now, updated_at=now))
        session.add(NoteVersion(note_id=note_id, revision=1, title="Frühere Frage",
                                relative_path=target.relative_to(settings.data_dir).as_posix(),
                                sha256=hashlib.sha256(content).hexdigest(), created_at=now))
        session.execute(text("INSERT INTO notes_fts(note_id, title, body) VALUES (:id, :title, :body)"),
                        {"id": note_id, "title": "Frühere Frage", "body": body})
        session.commit()
        return {"id": note_id}


@app.post("/api/__test__/document-date-proposal")
def document_date_proposal_for_browser_test():
    with app.state.factory() as session:
        resource = session.scalar(select(Resource).where(Resource.name == "04_Schleifen.pdf"))
        if resource is None:
            raise RuntimeError("Fake-Dokument fehlt")
        version = session.scalar(select(ResourceVersion).where(ResourceVersion.resource_id == resource.id)
                                 .order_by(ResourceVersion.version.desc()))
        if version is None:
            raise RuntimeError("Fake-Version fehlt")
        candidate = (utc_now() + timedelta(days=30)).date()
        proposal = DeadlineProposal(resource_id=resource.id, version_id=version.id, page_number=2,
                                    quote=f"Anmeldefrist bis zum {candidate:%d.%m.%Y}",
                                    candidate_date=candidate, status="proposed", created_at=utc_now())
        session.add(proposal)
        session.commit()
        return {"id": proposal.id}


@app.post("/api/__test__/budget-deadline")
def budget_deadline_for_browser_test():
    with app.state.factory() as session:
        deadline = Deadline(source="fixture", source_key=f"budget-{uuid4().hex}", course_id=1,
                            assignment_id=None, title="Planbare Übung",
                            due_at=utc_now() + timedelta(days=7), url=None)
        session.add(deadline)
        session.commit()
        return {"id": deadline.id}


@app.post("/api/__test__/submission-status")
def submission_status_for_browser_test(submitted: bool = False):
    global fake_submission_status
    fake_submission_status = "submitted" if submitted else "draft"
    return {"status": fake_submission_status}


@app.post("/api/__test__/submission-write-reset")
def reset_submission_write_for_browser_test():
    global fake_submission_status, uploaded_submission_name, uploaded_submission_content
    fake_submission_status = "draft"
    uploaded_submission_name = None
    uploaded_submission_content = None
    with app.state.factory() as session:
        session.execute(delete(SubmissionAttempt))
        session.commit()
    return {"status": fake_submission_status}


@app.post("/api/__test__/conflict-package")
def conflict_package_for_browser_test():
    host.assert_writer()
    package_id = f"{datetime.now(UTC):%Y%m%dT%H%M%S%fZ}-{uuid4().hex[:8]}"
    folder = ensure_directory(settings.data_dir, "host", "conflicts", package_id)
    action = {
        "kind": "intent", "id": uuid4().hex, "recorded_at": datetime.now(UTC).isoformat(),
        "method": "PUT", "path": "/api/ai/access/category/course_material",
        "body": {"allowed": False}, "body_state": "erfasst", "outcome": "erfolgreich",
    }
    manual_action = {
        "kind": "intent", "id": uuid4().hex, "recorded_at": datetime.now(UTC).isoformat(),
        "method": "SYSTEM", "path": "/internal/backup", "body": None,
        "body_state": "ohne JSON", "outcome": "erfolgreich",
    }
    actions = write_once(folder / "actions.json", json.dumps([action, manual_action]).encode("utf-8"))
    database_path = folder / "database.sqlite3"
    with sqlite3.connect(settings.local_dir / "test.sqlite3") as source, sqlite3.connect(database_path) as target:
        source.backup(target)
    manifest = {
        "package_id": package_id, "detected_at": datetime.now(UTC).isoformat(),
        "owned": {"device_name": "Test-PC"}, "winner": {"device_name": "Test-Mac"},
        "action_count": 2, "dashboard_actions_complete": True,
        "database_warning": None, "database_sha256": hashlib.sha256(database_path.read_bytes()).hexdigest(),
        "actions_sha256": hashlib.sha256(actions.read_bytes()).hexdigest(),
    }
    write_once(folder / "manifest.json", json.dumps(manifest).encode("utf-8"))
    return {"package_id": package_id, "action_id": action["id"], "manual_action_id": manual_action["id"]}


def capture_notifications():
    sent = []

    def fake_push(request):
        title = str(make_header(decode_header(request.headers["Title"])))
        sent.append({"title": title, "body": request.content.decode()})
        return httpx.Response(200)

    notifier_settings = settings.model_copy(update={"ntfy_url": "https://notify.example.invalid/ntfy"})
    notifier = Notifier(notifier_settings, host, client=httpx.Client(transport=httpx.MockTransport(fake_push)))
    return notifier, sent


@app.post("/api/__test__/doctor-credential-file")
def doctor_credential_file_for_browser_test(insecure: bool = False):
    doctor_module._credential_file = (
        lambda: Diagnostic("warning", ".env-Datei ist für andere Benutzer lesbar")
    ) if insecure else original_credential_file_check
    return {"insecure": insecure}


@app.post("/api/__test__/syncthing-conflict")
@app.get("/__test__/syncthing-conflict-reset")
def syncthing_conflict_for_browser_test(present: bool = False):
    restored = None
    if present:
        syncthing_conflict_file.write_text("erfundener Konflikt", encoding="utf-8")
    else:
        syncthing_conflict_file.unlink(missing_ok=True)
        # The conflict check revokes the writer for safety. Restore only this
        # isolated fake host after the browser scenario has removed its file.
        host.close()
        restored = host.start()
    return {"present": present, "restored": restored, "is_writer": host.is_writer}


@app.post("/api/__test__/upcoming-assignment")
def upcoming_assignment_for_browser_test():
    due_at = utc_now() + timedelta(days=3)
    with app.state.factory() as session:
        assignment = session.scalar(select(Assignment).where(Assignment.title == "Schleifen üben"))
        assert assignment is not None
        assignment.due_at = due_at
        for deadline in session.scalars(select(Deadline).where(Deadline.title == assignment.title)):
            deadline.due_at = due_at
        session.commit()
    return {"due_at": due_at.isoformat()}


@app.post("/api/__test__/deadline-reminder")
def deadline_reminder_for_browser_test():
    due_date = datetime.now(settings.timezone).date() + timedelta(days=3)
    due_at = datetime.combine(due_date, time(12), tzinfo=settings.timezone).astimezone(UTC)
    with app.state.factory() as session:
        existing = session.scalar(select(Deadline).where(Deadline.source == "browser-fixture",
                                                         Deadline.source_key == "reminder"))
        if existing is None:
            session.add(Deadline(source="browser-fixture", source_key="reminder",
                                 title="Abgabe aus Fake-Server", due_at=due_at))
            session.commit()
    notifier, sent = capture_notifications()
    scheduler = SyncScheduler(SyncEngine(app.state.factory, settings, [], notifier, host), settings, host)
    scheduler._warn_deadlines()
    return {"sent": sent}


@app.post("/api/__test__/expire-token")
def expire_token_for_browser_test():
    notifier, sent = capture_notifications()
    previous = fixture.responses["core_webservice_get_site_info"]
    fixture.responses["core_webservice_get_site_info"] = {
        "exception": "moodle_exception", "errorcode": "invalidtoken", "message": "Invalid token",
    }
    try:
        run = SyncEngine(app.state.factory, settings, [fixture.connector], notifier, host).run()
    finally:
        fixture.responses["core_webservice_get_site_info"] = previous
    return {"status": run.status, "sent": sent}


@app.post("/api/__test__/file-token-error")
def file_token_error_for_browser_test():
    notifier, sent = capture_notifications()
    file_path = "/webservice/pluginfile.php/41/Tokenfehler.pdf"
    response = b'{"exception":"moodle_exception","errorcode":"invalidtoken","message":"Invalid token"}'
    module = {"id": 541, "modname": "resource", "name": "Tokenfehler.pdf", "contents": [
        {"type": "file", "filename": "Tokenfehler.pdf",
         "fileurl": f"https://moodle.example.invalid{file_path}",
         "timemodified": 1727000001, "filesize": len(response)},
    ]}
    modules = fixture.contents[101][0]["modules"]
    modules.append(module)
    fixture.download_responses[file_path] = httpx.Response(
        200, content=response, headers={"content-type": "application/json"},
    )
    try:
        run = SyncEngine(app.state.factory, settings, [fixture.connector], notifier, host).run()
    finally:
        modules.remove(module)
        del fixture.download_responses[file_path]
    return {"status": run.status, "sent": sent}


@app.post("/api/__test__/queue-file-formats")
def queue_file_formats_for_browser_test():
    samples = sample_files()
    modules = fixture.contents[101][0]["modules"]
    for offset, (filename, data) in enumerate(samples.items(), start=600):
        source = f"/webservice/pluginfile.php/60/{filename}"
        fixture.files[source] = data
        if any(module.get("id") == offset for module in modules):
            continue
        modules.append({
            "id": offset, "modname": "resource", "name": filename,
            "contents": [{"type": "file", "filename": filename,
                          "fileurl": f"https://moodle.example.invalid{source}",
                          "timemodified": 1727000000, "filesize": len(data)}],
        })
    return {"sha256": {name: hashlib.sha256(data).hexdigest() for name, data in samples.items()},
            "download_count": fixture.download_count}


@app.get("/api/__test__/download-count")
def download_count_for_browser_test():
    return {"download_count": fixture.download_count}


@app.post("/api/__test__/queue-redirected-file")
def queue_redirected_file_for_browser_test():
    filename = "Umgeleitete_Vorlesung.pdf"
    source = f"/webservice/pluginfile.php/61/start/{filename}"
    target = f"/webservice/pluginfile.php/61/target/{filename}"
    fixture.download_responses[source] = httpx.Response(302, headers={"location": target})
    fixture.files[target] = pdf_data
    modules = fixture.contents[101][0]["modules"]
    if not any(module.get("id") == 610 for module in modules):
        modules.append({
            "id": 610, "modname": "resource", "name": filename,
            "contents": [{"type": "file", "filename": filename,
                          "fileurl": f"https://moodle.example.invalid{source}",
                          "timemodified": 1727000000, "filesize": len(pdf_data)}],
        })
    return {"sha256": hashlib.sha256(pdf_data).hexdigest(), "download_count": fixture.download_count}


@app.post("/api/__test__/broken-course-list")
def broken_course_list_for_browser_test():
    notifier, sent = capture_notifications()
    previous = fixture.responses["core_enrol_get_users_courses"]
    fixture.responses["core_enrol_get_users_courses"] = {"invalid": "course-list"}
    try:
        run = SyncEngine(app.state.factory, settings, [fixture.connector], notifier, host).run()
    finally:
        fixture.responses["core_enrol_get_users_courses"] = previous
    return {"status": run.status, "sent": sent}


@app.post("/api/__test__/wake-after-missed-sync")
def wake_after_missed_sync_for_browser_test():
    configured = settings.model_copy(update={"quiet_hours_start": None, "quiet_hours_end": None})
    scheduler = SyncScheduler(app.state.sync_engine, configured, host)
    completed = Event()
    scheduler.scheduler.add_listener(lambda _event: completed.set(), EVENT_JOB_EXECUTED)
    with app.state.factory() as session:
        before = session.scalar(select(SyncRun.id).order_by(SyncRun.id.desc()).limit(1))
    scheduler.scheduler.start(paused=True)
    try:
        scheduler._schedule_at(datetime.now(configured.timezone) - timedelta(hours=2))
        scheduler.scheduler.resume()
        if not completed.wait(10):
            raise RuntimeError("Versäumter Sync wurde nicht nachgeholt")
    finally:
        scheduler.stop()
    with app.state.factory() as session:
        latest = session.scalar(select(SyncRun).order_by(SyncRun.id.desc()).limit(1))
    return {"previous_id": before, "latest_id": latest.id, "status": latest.status}


@app.post("/api/__test__/backup-status")
def backup_status_for_browser_test(fail: bool = False):
    status_dir = settings.data_dir / "backup-status"
    status_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(UTC)
    status = {"status": "error" if fail else "success", "finished_at": timestamp.isoformat(),
              "message": "Backup fehlgeschlagen" if fail else "Test-Backup geprüft"}
    (status_dir / f"{timestamp:%Y%m%dT%H%M%S%fZ}.json").write_text(json.dumps(status), encoding="utf-8")
    return status


@app.post("/api/__test__/snapshot-status")
def snapshot_status_for_browser_test(fail: bool = False):
    if not fail:
        return asdict(create_snapshot(host))
    lease = host.current()
    assert lease is not None, "Test-Route benötigt einen bereits vorhandenen Staffelstab"
    now = datetime.now(UTC)
    filename = f"snapshot_{lease.generation:08d}_{now:%Y%m%dT%H%M%S%fZ}_{uuid4().hex[:8]}.sqlite3"
    invalid = SnapshotManifest(filename, lease.generation, now.isoformat(), "0" * 64, lease.device_id)
    path = settings.data_dir / "host" / "snapshots" / f"{filename}.json"
    with path.open("x", encoding="utf-8") as output:
        json.dump(asdict(invalid), output)
    return asdict(invalid)


@app.post("/api/__test__/change-file")
def change_file_for_browser_test():
    source = fixture.contents[101][0]["modules"][0]["contents"][0]
    source["timemodified"] += 1
    document = fitz.open(stream=pdf_data, filetype="pdf")
    document.new_page().insert_text((72, 72), "Aktualisierte Vorlesung")
    changed_pdf = document.tobytes()
    document.close()
    fixture.files["/webservice/pluginfile.php/10/04_Schleifen.pdf"] = changed_pdf
    source["filesize"] = len(changed_pdf)
    before_downloads = fixture.download_count
    run = app.state.sync_engine.run()
    return {"status": run.status, "new_downloads": fixture.download_count - before_downloads}


@app.post("/api/__test__/remove-file")
def remove_file_for_browser_test():
    fixture.contents[101][0]["modules"][0]["contents"] = []
    before_downloads = fixture.download_count
    run = app.state.sync_engine.run()
    return {"status": run.status, "new_downloads": fixture.download_count - before_downloads}


@app.post("/api/__test__/unsafe-urls")
def unsafe_urls_for_browser_test():
    names = ("mod_assign_get_assignments", "mod_forum_get_forum_discussions",
             "core_calendar_get_action_events_by_timesort")
    original = {name: copy.deepcopy(fixture.responses[name]) for name in names}
    fixture.responses["mod_assign_get_assignments"]["courses"][0]["assignments"][0]["url"] = (
        "https://moodle.example.invalid/mod/assign/view.php?id=701&token=example-value"
    )
    fixture.responses["mod_forum_get_forum_discussions"]["discussions"][0]["url"] = "javascript:alert(1)"
    fixture.responses["core_calendar_get_action_events_by_timesort"]["events"][0]["url"] = (
        "https://moodle.example.invalid/calendar/view.php?id=901&wstoken=example-value"
    )
    try:
        run = app.state.sync_engine.run()
    finally:
        fixture.responses.update(original)
    return {"status": run.status}


@app.post("/api/__test__/foreign-group")
def foreign_group_for_browser_test(new_section: bool = False):
    section = ({"id": 303, "name": "Gruppe Weber", "modules": []} if new_section else
               next(item for item in fixture.contents[101] if item["id"] == 302))
    if new_section:
        fixture.contents[101].append(section)
    filename = "Noch_Fremd.pdf" if new_section else "Fremdes_Material.pdf"
    file_path = f"/webservice/pluginfile.php/30/{filename}"
    fixture.files[file_path] = pdf_data
    modules = section["modules"]
    assert isinstance(modules, list), "Fixture-Abschnitt braucht eine Modulliste"
    modules.append({"id": 508 if new_section else 507, "modname": "resource", "name": filename,
                    "contents": [{"type": "file", "filename": filename,
                                  "fileurl": f"https://moodle.example.invalid{file_path}",
                                  "timemodified": 1727000000, "filesize": len(pdf_data)}]})
    with app.state.factory() as session:
        before_notifications = len(session.scalars(select(Notification).where(
            Notification.title.like("%neue Uni-Datei%"))).all())
    run = app.state.sync_engine.run()
    with app.state.factory() as session:
        after_notifications = len(session.scalars(select(Notification).where(
            Notification.title.like("%neue Uni-Datei%"))).all())
    return {"status": run.status, "download_count": fixture.download_count,
            "new_file_notifications": after_notifications - before_notifications}


@app.post("/api/__test__/sync-gap")
def sync_gap_for_browser_test(recover: bool = False):
    if recover:
        run = app.state.sync_engine.run()
        return {"status": run.status}
    old = utc_now() - timedelta(hours=38, minutes=5)
    with app.state.factory() as session:
        for run in session.scalars(select(SyncRun)).all():
            if run.status == "success" or (run.status == "warning" and run.course_count > 0):
                run.started_at = old
                run.finished_at = old
        session.add(SyncRun(status="error", finished_at=utc_now()))
        session.commit()
    notifier, sent = capture_notifications()
    scheduler = SyncScheduler(SyncEngine(app.state.factory, settings, [], notifier, host), settings, host)
    scheduler._warn_stale_sync()
    scheduler._warn_stale_sync()
    return {"sent": sent}


@app.post("/api/__test__/new-file-push")
def new_file_push_for_browser_test():
    filename = "Neue_Hinweise.pdf"
    file_path = f"/webservice/pluginfile.php/40/{filename}"
    fixture.files[file_path] = pdf_data
    fixture.contents[101][0]["modules"].append({
        "id": 509, "modname": "resource", "name": filename,
        "contents": [{"type": "file", "filename": filename,
                      "fileurl": f"https://moodle.example.invalid{file_path}",
                      "timemodified": 1727000000, "filesize": len(pdf_data)}],
    })
    notifier, sent = capture_notifications()
    engine = SyncEngine(app.state.factory, settings, [fixture.connector], notifier, host)
    before_downloads = fixture.download_count
    first = engine.run()
    first_downloads = fixture.download_count
    second = engine.run()
    return {"first_status": first.status, "second_status": second.status,
            "first_downloads": first_downloads - before_downloads,
            "second_downloads": fixture.download_count - first_downloads, "sent": sent}


@app.post("/api/__test__/queue-file-for-palette")
def queue_file_for_palette_test():
    filename = "Befehlspalette_Aktuell.pdf"
    file_path = f"/webservice/pluginfile.php/41/{filename}"
    fixture.files[file_path] = pdf_data
    fixture.contents[101][0]["modules"].append({
        "id": 510, "modname": "resource", "name": filename,
        "contents": [{"type": "file", "filename": filename,
                      "fileurl": f"https://moodle.example.invalid{file_path}",
                      "timemodified": 1727000001, "filesize": len(pdf_data)}],
    })
    return {"queued": filename}


@app.post("/api/__test__/moodle-assignment-diagnostics")
def moodle_assignment_diagnostics_for_browser_test(expected_access: bool = False):
    previous = copy.deepcopy(fixture.responses["mod_assign_get_assignments"])
    contents = copy.deepcopy(fixture.contents[101])
    try:
        fixture.contents[101].append({
            "id": 9000, "uservisible": False,
            "modules": [{"id": 9001, "modname": "assign", "name": "Fictional hidden assignment"}],
        })
        fixture.responses["mod_assign_get_assignments"] = {
            "courses": [], "warnings": [
                {"item": "module", "itemid": 9001, "warningcode": "1", "message": "fixture-private-message"},
                {"item": "course", "itemid": 202, "warningcode": "2", "message": "fixture-private-message"},
            ],
        }
        if expected_access:
            fixture.responses["mod_assign_get_assignments"]["courses"] = copy.deepcopy(previous["courses"])
            fixture.responses["mod_assign_get_assignments"]["warnings"] = [
                {"item": "module", "itemid": 9001, "warningcode": "1", "message": "fixture-private-message"},
            ]
            for course in fixture.responses["mod_assign_get_assignments"]["courses"]:
                for assignment in course["assignments"]:
                    assignment["cmid"] = 9002
            fixture.contents[101][0]["modules"].append({"id": 9002, "modname": "assign"})
        result = app.state.sync_engine.run()
        return {"status": result.status}
    finally:
        fixture.responses["mod_assign_get_assignments"] = previous
        fixture.contents[101] = contents


@app.post("/api/__test__/moodle-update-hint")
def moodle_update_hint_for_browser_test():
    filename = "Versionshinweis.txt"
    file_path = f"/webservice/pluginfile.php/50/{filename}"
    fixture.files[file_path] = b"Version A\n"
    fixture.contents[101][0]["modules"].append({
        "id": 510, "modname": "resource", "name": filename,
        "contents": [{"type": "file", "filename": filename,
                      "fileurl": f"https://moodle.example.invalid{file_path}",
                      "timemodified": 1727000000, "filesize": len(fixture.files[file_path])}],
    })
    before_downloads = fixture.download_count
    first = app.state.sync_engine.run()
    first_downloads = fixture.download_count
    fixture.files[file_path] = b"Version B\n"
    fixture.responses["core_course_get_updates_since"] = {
        "instances": [{"contextlevel": "module", "id": 510,
                       "updates": [{"name": "fileareas", "timeupdated": 1727000001}]}],
        "warnings": [],
    }
    try:
        second = app.state.sync_engine.run()
    finally:
        fixture.responses["core_course_get_updates_since"] = {"instances": [], "warnings": []}
    second_downloads = fixture.download_count
    third = app.state.sync_engine.run()
    return {"statuses": [first.status, second.status, third.status],
            "downloads": [first_downloads - before_downloads,
                          second_downloads - first_downloads,
                          fixture.download_count - second_downloads]}


@app.post("/api/__test__/remove-assignment")
def remove_assignment_for_browser_test():
    assignments = fixture.responses["mod_assign_get_assignments"]["courses"][0]["assignments"]
    fixture.responses["mod_assign_get_assignments"]["courses"][0]["assignments"] = []
    try:
        run = app.state.sync_engine.run()
    finally:
        fixture.responses["mod_assign_get_assignments"]["courses"][0]["assignments"] = assignments
    return {"status": run.status}


@app.post("/api/__test__/schedule-exam-in-5-days")
def schedule_exam_for_browser_test():
    with app.state.factory() as session:
        session.add(Exam(course_id=1, source_key="fixture-exam", title="Testklausur",
                         exam_at=utc_now() + timedelta(days=5)))
        session.commit()
    return {"status": "success"}


@app.post("/api/__test__/accepted-study-block-today")
def create_accepted_study_block_today():
    """An already-accepted study block starting right now, for the focus-timer UI.

    The normal path (schedule an exam, get a rule-based suggestion, accept
    it) almost never lands on today's date -- ``suggest_blocks`` schedules
    the day before the exam, which is only "today" for an exam tomorrow, an
    edge case sensitive to what time of day the suite happens to run. Since
    the focus timer only cares that a block is accepted and starts today,
    this inserts one directly.
    """
    with app.state.factory() as session:
        course = session.get(Course, 1)
        # A fresh topic name per call: several tests within the same spec file share this
        # long-lived server without an in-between reset, and each needs its own,
        # unambiguous row in the focus-timer list.
        topic = resolve_topic(session, course.id, f"Fokus-Test-Thema-{uuid4().hex[:8]}")
        exam = Exam(course_id=course.id, source_key=f"fixture-focus-exam-{uuid4().hex}",
                   title="Fokus-Testklausur", exam_at=utc_now() + timedelta(days=2))
        session.add(exam)
        session.flush()
        block = StudyBlockSuggestion(exam_id=exam.id, topic_id=topic.id, duration_minutes=1,
                                     starts_at=utc_now(), status="accepted", created_at=utc_now())
        session.add(block)
        session.commit()
        return {"id": block.id, "topic": topic.name, "course_name": course.name,
                "duration_minutes": block.duration_minutes}


class _StubAiAdapter:
    """Deterministic stand-in for a real provider, used by ``ai-adapter/enable`` below.

    Reads the resource/version the prompt already names in its manifest
    (built by ``build_ai_package``) instead of hardcoding one, so it stays
    correct for whichever course a test targets. Which schema file the
    prompt embeds (``card_schema.json`` vs. ``exercise_schema.json``) tells
    it which shape of proposal to return.
    """

    def propose(self, prompt: str) -> AdapterResult:
        resource_match = re.search(r'"resource_id":\s*(\d+)', prompt)
        version_match = re.search(r'"version":\s*(\d+)', prompt)
        assert resource_match is not None
        assert version_match is not None
        resource_id = int(resource_match.group(1))
        version = int(version_match.group(1))
        source = {"resource_id": resource_id, "version": version}
        generated_at = utc_now().isoformat()
        if "card_schema.json" in prompt:
            body = {"schema_version": 1, "generated_by": "stub-adapter", "generated_at": generated_at,
                    "source": source, "cards": [{"card_type": "basic", "front": "Stub-Frage",
                    "back": "Stub-Antwort", "course": "Stub-Kurs", "topic": "Stub-Thema", "confidence": "medium"}]}
        else:
            body = {"schema_version": 1, "generated_by": "stub-adapter", "generated_at": generated_at,
                    "source": source, "exercises": [{"question": "Stub-Frage?", "topic": "Stub-Thema",
                    "difficulty": "easy", "sample_solution": "Stub-Lösung", "solution_steps": "Stub-Schritt",
                    "source_page": 1}]}
        return AdapterResult(model="stub-model", tokens_used=42, payload=json.dumps(body, ensure_ascii=False))


class _FakeModelList:
    @staticmethod
    def list(**_kwargs):
        return SimpleNamespace(data=[SimpleNamespace(id="fixture-model-new", display_name="Fixture Modell Neu"),
                                     SimpleNamespace(id="fixture-model-old", display_name="Fixture Modell Alt")])


# The provider flow never reaches the network or the developer's keychain from the browser suite.
_fixture_vault: dict[tuple[str, str], str] = {}


def _fixture_save(connector: str, alias: str, value: str) -> bool:
    _fixture_vault[(connector, alias)] = value
    return True


ai_provider_service._client = lambda _key: SimpleNamespace(models=_FakeModelList())  # type: ignore[assignment]
ai_provider_repository.load_connector_secret = lambda connector, alias: _fixture_vault.get((connector, alias), "")  # type: ignore[assignment]
ai_provider_repository.save_connector_secret = _fixture_save  # type: ignore[assignment]
ai_provider_repository.delete_connector_secret = lambda connector, alias: _fixture_vault.pop((connector, alias), None)  # type: ignore[assignment]


@app.post("/api/__test__/ai-adapter/enable")
def enable_fake_ai_adapter():
    """Simulate an enabled, configured KI-Adapter with a stub provider (no real network call)."""
    settings.ai_adapter_enabled = True
    settings.ai_adapter_api_key = "fixture-adapter-key"  # pragma: allowlist secret
    settings.ai_adapter_model = "fixture-model-new"
    ai_api_module.select_adapter = lambda _settings: _StubAiAdapter()
    return {"status": "enabled"}


@app.post("/api/__test__/ai-adapter/disable")
def disable_fake_ai_adapter():
    settings.ai_adapter_enabled = False
    settings.ai_adapter_api_key = ""
    settings.ai_adapter_model = ""
    _fixture_vault.clear()
    ai_api_module.select_adapter = _original_select_adapter
    return {"status": "disabled"}


@app.post("/api/__test__/reset")
def reset_for_browser_test():
    """Restore the fake server to its just-started baseline for the next spec.

    Every previous ``/api/__test__/*`` route above resets one narrow slice of
    state; specs that don't call the right one leak data into whichever spec
    runs next on this long-lived fake server (stray ``submission_attempts``,
    duplicated approved exercises, cumulative fixture edits). This route
    resets everything in one call:

    1. Restore the SQLite database file itself, byte-for-byte, from the
       verified snapshot taken right after startup finished its own sync —
       every table, including the FTS5 search indexes, goes back to exactly
       that state. This intentionally does not delete rows and re-run the
       sync: the immutable, content-addressed files already on disk under
       data/uni/ are untouched by a reset, and re-downloading into an emptied
       database would re-create their per-version sidecar JSON with a new
       timestamp, which the write-once file store correctly refuses to write
       over the still-present original.
    2. Restore the in-memory Moodle ``fixture`` (contents/responses/files and
       the download counter) to the deep copy taken at that same moment, and
       reset the fake submission-draft globals.
    3. Clear the DATA_DIR subfolders tests write into (annotation journals,
       drop, the AI inbox/outbox, the submissions archive, versioned note and
       course-scratchpad Markdown files), then restore the one file that was
       seeded before the app ever started. Notes and the course scratchpad
       are indexed by a DB revision number that this reset rewinds to 0; the
       write-once file store refuses to overwrite a same-named file left
       behind by an earlier spec, so its folder must be cleared too or the
       very next write for that course collides and fails with a 500.

    A Playwright global setup or per-spec ``beforeEach`` should ``POST`` here
    before every spec (or every test) that does not need to share state with
    a previous one, instead of relying on this one long-lived server to stay
    untouched.
    """
    global fake_submission_status, uploaded_submission_name, uploaded_submission_content

    from app.anki.collection import close_study_collection

    close_study_collection(settings)
    # Only this process's TemporaryDirectory is ever reset. Remove conflicts
    # before restoring its fictional coordinator, including lease-loss scenarios.
    for path in (settings.data_dir / "host").glob("*.sync-conflict-*"):
        path.unlink(missing_ok=True)
    host.close()
    host.lease_path.write_bytes(_lease_baseline)
    host.start()
    _disarm_review_fault()
    _reset_study_fixture()
    engine = app.state.factory.kw["bind"]
    engine.dispose()
    with sqlite3.connect(_reset_baseline_db) as baseline_connection, sqlite3.connect(_live_db_path) as live_connection:
        baseline_connection.backup(live_connection)

    fixture.contents = copy.deepcopy(_fixture_baseline["contents"])
    fixture.responses = copy.deepcopy(_fixture_baseline["responses"])
    fixture.files = copy.deepcopy(_fixture_baseline["files"])
    fixture.download_count = _fixture_baseline["download_count"]
    fixture.download_responses = {}
    fixture.calls = []
    fixture.site_info_called_at = None
    fake_submission_status = "draft"
    uploaded_submission_name = None
    uploaded_submission_content = None
    settings.ai_adapter_enabled = False
    settings.ai_adapter_api_key = ""
    settings.ai_adapter_model = ""
    _fixture_vault.clear()
    ai_api_module.select_adapter = _original_select_adapter
    app.state.corpus_cache.clear()
    doctor_module._credential_file = original_credential_file_check
    fake_forwarding.configuration.save({})
    fake_forwarding.state.save({})
    for name in ("new", "cur"):
        folder = fake_forwarding.mailbox / name
        if folder.is_dir():
            for path in folder.iterdir():
                if path.is_file() and not path.is_symlink():
                    path.unlink()
    for name in ("addon-events.jsonl", "addon-catalog.jsonl"):
        (settings.local_dir / name).unlink(missing_ok=True)

    from app.integrations.storage import ConnectionStore
    ConnectionStore(settings.local_dir, 'timetable').save({})
    ConnectionStore(settings.local_dir, 'timetable-selection').save({})

    for relative in _RESET_DATA_SUBFOLDERS:
        folder = settings.data_dir / Path(relative)
        if folder.is_symlink() or not folder.resolve().is_relative_to(_data_dir.resolve()):
            continue
        if folder.is_dir():
            shutil.rmtree(folder, ignore_errors=True)
    snapshots = settings.data_dir / "host/snapshots"
    snapshots.mkdir(parents=True, exist_ok=True)
    for name, content in _snapshot_baseline.items():
        (snapshots / name).write_bytes(content)
    drop = settings.data_dir / "drop"
    drop.mkdir(parents=True, exist_ok=True)
    (drop / "Eigene_Frage.txt").write_text(_seed_drop_question, encoding="utf-8")

    return {"status": "ok"}


@app.post("/api/__test__/exam-impact-plan")
def create_exam_impact_plan():
    """Create a changed exam and an accepted plan in the isolated fake database."""
    now = utc_now()
    previous_exam = now + timedelta(days=14)
    current_exam = now + timedelta(days=12)
    suffix = uuid4().hex
    with app.state.factory() as session:
        topic = resolve_topic(session, 1, f"Planänderung {suffix[:6]}")
        exam = Exam(course_id=1, source_key=f"fixture-impact-{suffix}", title="Synthetische Terminänderung",
                    exam_at=current_exam)
        session.add(exam)
        session.flush()
        cycle = StudyCycle(exam_id=exam.id, phase="practice", starts_at=now - timedelta(days=1),
                           ends_at=previous_exam - timedelta(days=3), exam_at_snapshot=previous_exam,
                           topic_ids=[topic.id], status="accepted")
        review = StudyCycle(exam_id=exam.id, phase="review", starts_at=previous_exam - timedelta(days=3),
                            ends_at=previous_exam, exam_at_snapshot=previous_exam, topic_ids=[topic.id], status="accepted")
        block = StudyBlockSuggestion(exam_id=exam.id, topic_id=topic.id, starts_at=previous_exam - timedelta(days=1),
                                     duration_minutes=45, status="accepted")
        session.add_all([cycle, review, block])
        session.commit()
        return {"cycle_id": cycle.id, "block_id": block.id, "exam_id": exam.id,
                "old_exam_at": previous_exam.isoformat(), "current_exam_at": current_exam.isoformat()}
